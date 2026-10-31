"""Fetch, build and validate the sources a dataset is built from, as `sources.toml` declares them.

Raw files stay in data/raw/ and outputs in data/processed/. Each fetch leaves a
`<file>.fetch.json` record beside its raw file: the URL, when it was retrieved, the SHA-256
of the bytes, the vintage the manifest declares, and the server's Last-Modified and ETag
headers. A later fetch sends those headers back, and when the server answers 304 Not
Modified it keeps the raw file and its record as they are.

Each fetch that brings bytes a source has not had before also appends its record to
data/fetches.jsonl, the history of every version of every source the pipeline has seen.

Paths are relative to the working directory, the project root.
"""

from __future__ import annotations

import contextlib
import csv
import dataclasses
import errno
import functools
import gzip
import hashlib
import http.client
import io
import ipaddress
import itertools
import json
import logging
import os
import re
import shutil
import socket
import sys
import tempfile
import time
import tomllib
import urllib.error
import urllib.request
import urllib.response
import zipfile
import zlib
from collections import Counter
from collections.abc import Callable, Collection, Generator, Iterable, Iterator
from datetime import UTC, date, datetime
from email.utils import parsedate_to_datetime
from http import HTTPStatus
from importlib.metadata import version
from pathlib import Path, PurePosixPath
from typing import IO, TYPE_CHECKING, Any
from urllib.parse import urljoin, urlsplit

if TYPE_CHECKING:
    from email.message import Message

if sys.platform != "win32":  # noqa: PLR2004
    import fcntl

MANIFEST = Path("sources.toml")
RAW = Path("data/raw")
PROCESSED = Path("data/processed")
FETCH_LOG = Path("data/fetches.jsonl")
# Written by build beside the outputs.
DATAPACKAGE = "datapackage.json"
_DATAPACKAGE_PROFILE = "https://datapackage.org/profiles/2.0/datapackage.json"
# A manifest's first line names it, for an editor to check the file. The copy at this
# version's tag, not the $id on main: main can accept or refuse keys this version doesn't.
_SCHEMA_URL = (
    f"https://raw.githubusercontent.com/simonvanlierde/sourceledger/v{version('sourceledger')}"
    "/src/sourceledger/sources.schema.json"
)
_STARTER = f"""#:schema {_SCHEMA_URL}

# One [[source]] table per input. Uncomment the example and point it at your own source.
# Every key: https://github.com/simonvanlierde/sourceledger#sourcestoml

# [[source]]
# id = "co2"
# status = "implemented"
# url = "https://raw.githubusercontent.com/datasets/co2-ppm/main/data/co2-annmean-mlo.csv"
# vintage = "main branch, unversioned"
# terms = "ODC-PDDL-1.0 (datasets/co2-ppm package; original data from NOAA GML)"
# terms_url = "https://opendatacommons.org/licenses/pddl/"
# license = "PDDL-1.0"
# redistribute = true
"""

IMPLEMENTED = "implemented"
UNIMPLEMENTED = "unimplemented"
FILE_FORMAT = "file"
# Bounds each connect and each read. A server that trickles bytes is stopped by _BUDGET_S.
_SOCKET_TIMEOUT_S = 30
# One source's download, every attempt and wait included, so a stalled publisher costs
# at most 5 minutes of a CI job unless timeout_s overrides it.
_BUDGET_S = 300
# Waits of 5, 10 and 20 s between four attempts ride out a brief publisher outage.
_ATTEMPTS = 4
_BACKOFF_S = 5
# A server asking for hours would otherwise stall the run.
_MAX_RETRY_AFTER_S = 120
# Statuses a publisher recovers from. 501 and 505 never will.
_RETRIED = frozenset({429, 500, 502, 503, 504})
_CHUNK_BYTES = 1024 * 1024
# Hashing a file already on disk: small enough that its peak memory stays small, large
# enough that a small file is one read.
_HASH_BLOCK_BYTES = 64 * 1024
# A network error or timeout, or a connection dropped or garbled mid-response.
_TRANSIENT = (OSError, http.client.HTTPException)
# A raw file or fetch record that is missing, edited or unreadable.
_UNUSABLE = (OSError, ValueError)
# A runaway or wrong URL fails the fetch instead of filling the disk. A big source sets
# its own `max_bytes` in sources.toml.
_MAX_BYTES = 200 * 1024 * 1024
_SHA256 = re.compile(r"[0-9a-f]{64}")
_GZIP_MAGIC = b"\x1f\x8b"
_ENV_REF = re.compile(r"env:([A-Za-z_][A-Za-z0-9_]*)")
# Headers a manifest may not set: the fetch's own.
_RESERVED_HEADERS = frozenset({"user-agent", "accept-encoding", "if-none-match", "if-modified-since"})
# Sent again after a redirect to another scheme, host or port. Credentials from the manifest are not.
_SAFE_ON_REDIRECT = frozenset({"user-agent", "accept-encoding"})
# The Internet Archive's Save Page Now: a GET of this plus a URL asks it to snapshot that URL.
_ARCHIVE = "https://web.archive.org/save/"
_SNAPSHOT = "https://web.archive.org/web/"
# Save Page Now answers only once the snapshot is taken, 49 s for a GitHub page on 2026-10-20, so
# this bounds each connect and read of that request.
_ARCHIVE_TIMEOUT_S = 120

# Sent with every request, so a publisher can tell who is fetching. `main` takes the project's own.
USER_AGENT = f"sourceledger/{version('sourceledger')}"

_logger = logging.getLogger("sourceledger")

type Source = dict[str, Any]
# What the opener returns: an HTTP response, or a file:// one.
type _Response = http.client.HTTPResponse | urllib.response.addinfourl
# Reads one raw file as its columns and rows. `parse_source` is the default.
type Parser = Callable[[Source, bytes], tuple[list[str], list[dict[str, str]]]]


@dataclasses.dataclass(frozen=True, slots=True)
class FetchRecord:
    """What one download was, so every output row traces back to the bytes it came from."""

    source_id: str
    url: str
    retrieved_at: str
    sha256: str
    bytes: int
    vintage: str
    # Empty when the server sends none. Defaults keep records written before these fields readable.
    last_modified: str = ""
    etag: str = ""
    # Where the bytes came from after redirects. Empty when there were none.
    final_url: str = ""
    # The request headers taken from environment variables (names only, never values).
    env_headers: list[str] = dataclasses.field(default_factory=list)
    # The Internet Archive's snapshot of the URL, for a source with `archive = true`. Set by
    # the fetch that first brought these bytes; a later download of the same bytes takes it
    # from their line in the log. Empty when the source does not ask for one, or the archive
    # did not answer with one.
    archive_url: str = ""


class FetchError(OSError):
    """One or more sources failed to fetch. The others were fetched and recorded as usual."""

    def __init__(self, records: list[FetchRecord], failures: list[str]) -> None:
        super().__init__("\n".join(failures))
        self.records = records  # the sources that succeeded
        self.failures = failures  # one line per source that did not


def load_manifest(path: Path = MANIFEST) -> list[Source]:
    """The `[[source]]` tables of sources.toml, after checking ids, URLs, caps, timeouts and pins."""
    with path.open("rb") as fh:
        sources = tomllib.load(fh).get("source", [])
    if not isinstance(sources, list) or not all(isinstance(s, dict) and isinstance(s.get("id"), str) for s in sources):
        msg = "sources.toml: every [[source]] needs a string id"
        raise ValueError(msg)
    ids = [s["id"] for s in sources]
    # ASCII only: an id names files, and must mean the same file on every system.
    if bad := [i for i in ids if not re.fullmatch(r"[\w-]+", i, re.ASCII)]:
        msg = f"ids {bad} in sources.toml may hold only ASCII letters, digits, _ and -"
        raise ValueError(msg)
    if dupes := sorted(i for i, n in Counter(ids).items() if n > 1):
        msg = f"ids {dupes} appear more than once in sources.toml"
        raise ValueError(msg)
    if no_url := [s["id"] for s in implemented(sources) if not isinstance(s.get("url"), str)]:
        msg = f"implemented sources {no_url} in sources.toml have no url"
        raise ValueError(msg)
    # Not isinstance: bool is an int subclass, so `max_bytes = true` would pass as a 1-byte cap.
    caps = {s["id"]: s.get("max_bytes", _MAX_BYTES) for s in sources}
    if bad_cap := [i for i, cap in caps.items() if type(cap) is not int or cap < 1]:
        msg = f"sources {bad_cap} in sources.toml need max_bytes to be a positive whole number"
        raise ValueError(msg)
    if bad_archive := [s["id"] for s in sources if not isinstance(s.get("archive", False), bool)]:
        msg = f"sources {bad_archive} in sources.toml need archive to be true or false"
        raise ValueError(msg)
    timeouts = {s["id"]: s.get("timeout_s", _BUDGET_S) for s in sources}
    if bad_timeout := [
        sid
        for sid, timeout in timeouts.items()
        if type(timeout) not in (int, float) or not 0 < timeout <= sys.float_info.max
    ]:
        msg = f"sources {bad_timeout} in sources.toml need timeout_s to be a positive finite number"
        raise ValueError(msg)
    for source in sources:
        _check_pin(source)
        _check_vintage(source)
        _check_format(source)
        _check_headers(source)
    return sources


def init(path: Path = MANIFEST) -> None:
    """Write a starter sources.toml at `path`: the schema line and one commented example source.

    Raises FileExistsError when `path` exists, which is never overwritten.
    """
    try:
        with path.open("x", encoding="utf-8") as fh:
            fh.write(_STARTER)
    except FileExistsError as error:
        msg = f"{path} already exists; not overwriting it"
        raise FileExistsError(msg) from error


def _select(sources: list[Source], only: Collection[str] | None) -> list[Source]:
    """The sources whose ids are in `only`, in manifest order, or all of them when `only` is None.

    Raises ValueError for an id the manifest does not have.
    """
    if only is None:
        return sources
    if isinstance(only, str):  # a str is a Collection of its characters
        msg = f"only takes a list of ids, not the string {only!r}"
        raise TypeError(msg)
    if unknown := sorted(set(only) - {s["id"] for s in sources}):
        msg = f"ids {unknown} are not in sources.toml"
        raise ValueError(msg)
    return [s for s in sources if s["id"] in only]


def _check_format(source: Source) -> None:
    if (fmt := source.get("format")) is not None and fmt not in (*_PARSERS, FILE_FORMAT):
        msg = f"{source['id']}: format in sources.toml must be one of {sorted([*_PARSERS, FILE_FORMAT])}, not {fmt!r}"
        raise ValueError(msg)


def _check_headers(source: Source) -> None:
    """Each extra request header takes its value from an environment variable, never a literal, and goes over https."""
    headers = source.get("headers", {})
    if not isinstance(headers, dict):
        msg = f"{source['id']}: headers in sources.toml must be a table"
        raise ValueError(msg)  # noqa: TRY004  # a manifest error: the CLI reports ValueError as one line
    url = source.get("url")
    # The same test as the schema's pattern, so the two accept the same manifests.
    if source.get("headers") is not None and isinstance(url, str) and not url.lower().startswith("https:"):
        msg = f"{source['id']}: headers in sources.toml need an https url, so the values are never sent in clear text"
        raise ValueError(msg)
    for name, value in headers.items():
        if name.lower() in _RESERVED_HEADERS:
            msg = f"{source['id']}: sourceledger sets the {name} header itself"
            raise ValueError(msg)
        if not (isinstance(value, str) and _ENV_REF.fullmatch(value)):
            msg = (
                f'{source["id"]}: header {name} must name an environment variable, as "env:NAME"; '
                "never write a secret into sources.toml"
            )
            raise ValueError(msg)


def _check_pin(source: Source) -> None:
    """Normalize a `sha256` pin to lower case, as `sha256sum` prints it, and reject one that is not a SHA-256."""
    if (pin := source.get("sha256")) is None:
        return
    if not (isinstance(pin, str) and _SHA256.fullmatch(pin.lower())):
        msg = f"{source['id']}: sha256 in sources.toml must be 64 hexadecimal characters"
        raise ValueError(msg)
    source["sha256"] = pin.lower()


def _check_vintage(source: Source) -> None:
    """Refuse the old key `release`: kept as an unknown key, it would leave the vintage "not stated"."""
    if source.get("release") is not None:
        msg = f"{source['id']}: `release` in sources.toml is now `vintage`; rename it"
        raise ValueError(msg)


def _vintage(source: Source) -> str:
    # str(): a TOML date (`vintage = 2024-06-01`) loads as a datetime.date.
    return str(source.get("vintage", "not stated"))


def implemented(sources: list[Source]) -> list[Source]:
    """The sources the pipeline reads."""
    return [s for s in sources if s.get("status") == IMPLEMENTED]


def raw_name(source: Source) -> str:
    """The raw file's name: the source id plus the URL's extension, e.g. `example-co2.csv`."""
    return source["id"] + PurePosixPath(urlsplit(source["url"]).path).suffix


def _output_name(source: Source) -> str:
    """The built file's name; an unimplemented source may have no URL yet."""
    name = raw_name(source) if source.get("format") == FILE_FORMAT and source.get("url") else f"{source['id']}.csv"
    if name.lower() == DATAPACKAGE:
        msg = f"{DATAPACKAGE} is reserved for the package descriptor; choose another source id"
        raise ValueError(msg)
    return name


def _write_atomic(path: Path, data: bytes) -> None:
    """Write `data` to `path` as `_replacing` does."""
    with _replacing(path) as fh:
        fh.write(data)


@contextlib.contextmanager
def _replacing(path: Path) -> Generator[IO[bytes]]:
    """A new temp file beside `path` to write; on success it is flushed to disk and renamed over `path`.

    A reader never sees half a file. The temp file gets a fresh name, so two runs at once
    don't share it, and it is created rather than opened, so a symlink in its place is
    never followed. On an error the temp file is removed and `path` is left as it was.
    """
    fd, temp = tempfile.mkstemp(dir=path.parent, prefix=f".{path.name}.", suffix=".part")
    try:
        with os.fdopen(fd, "wb") as fh:
            yield fh
            fh.flush()
            os.fsync(fh.fileno())
        Path(temp).replace(path)
    except BaseException:
        Path(temp).unlink(missing_ok=True)
        raise


def _is_public(host: str) -> bool:
    """Whether every address `host` resolves to is on the public internet."""
    try:
        infos = socket.getaddrinfo(host, None)
    except OSError:
        return True  # unresolvable: the request itself will fail
    return all(ipaddress.ip_address(info[4][0]).is_global for info in infos)


class _InwardRedirectError(ValueError):
    """A redirect from a public host connected to a loopback, private or link-local address."""


class _NoRedirectInward(urllib.request.HTTPRedirectHandler):
    """Refuses a redirect from a public host to a loopback, private or link-local one.

    Otherwise a third-party server could have the fetch save an internal service's response
    under the manifest's URL. A source that lives on an internal host can still redirect
    within it. The redirect is only marked here; the connection checks the address it
    actually reached (`_PeerCheck`), so a name that resolves to another address the second
    time can't slip through.
    """

    def redirect_request(
        self,
        req: urllib.request.Request,
        fp: IO[bytes],
        code: int,
        msg: str,
        headers: http.client.HTTPMessage,
        newurl: str,
    ) -> urllib.request.Request | None:
        """Follow the redirect as urllib does, marking it public-only when it comes from a public host."""
        new = super().redirect_request(req, fp, code, msg, headers, newurl)
        if new is None:
            return None
        # Every later hop stays public-only, so a public host can't bounce through one that looks internal.
        if getattr(req, "public_only", False) or _is_public(urlsplit(req.full_url).hostname or ""):
            new.public_only = True  # ty: ignore[unresolved-attribute]  # read back by _PeerCheck
        # urllib copies every header to the new request. A token meant for one host must not
        # reach another, nor go out in clear text after https -> http, so a redirect that
        # changes the scheme, host or port keeps only the harmless ones.
        if urlsplit(newurl)[:2] != urlsplit(req.full_url)[:2]:
            new.headers = {k: v for k, v in new.headers.items() if k.lower() in _SAFE_ON_REDIRECT}
        return new


class _PeerCheck(http.client.HTTPConnection):
    """Once connected, refuses a public-only request whose socket reached a non-public address.

    Checked before anything is sent. Through a proxy the peer is the proxy, which resolves
    the name itself, so the check is skipped there.
    """

    def __init__(self, host: str, *, req: urllib.request.Request, **kwargs: Any) -> None:  # noqa: ANN401  # http.client's own arguments
        super().__init__(host, **kwargs)
        self.req = req

    def connect(self) -> None:
        """Connect as usual, then check the address the socket reached."""
        super().connect()
        target = urlsplit(self.req.full_url).hostname or ""
        if not getattr(self.req, "public_only", False) or self.host.lower() != target:
            return
        address = self.sock.getpeername()[0]
        if not ipaddress.ip_address(address).is_global:
            self.close()
            msg = f"refused a redirect from a public host to {self.req.full_url}, which is at {address}"
            raise _InwardRedirectError(msg)


class _HTTPSPeerCheck(_PeerCheck, http.client.HTTPSConnection):
    pass


class _HTTPHandler(urllib.request.HTTPHandler):
    def http_open(self, req: urllib.request.Request) -> http.client.HTTPResponse:
        return self.do_open(functools.partial(_PeerCheck, req=req), req)


class _HTTPSHandler(urllib.request.HTTPSHandler):
    def https_open(self, req: urllib.request.Request) -> http.client.HTTPResponse:
        return self.do_open(functools.partial(_HTTPSPeerCheck, req=req), req, context=self._context)  # ty: ignore[unresolved-attribute]  # set by HTTPSHandler.__init__


_OPENER = urllib.request.build_opener(_NoRedirectInward(), _HTTPHandler(), _HTTPSHandler())


def _open(request: urllib.request.Request, timeout: float) -> _Response:
    # NOTE: the opener follows file:// too, which the tests use for fixtures. The URLs come
    # from sources.toml, which is as trusted as the code.
    return _OPENER.open(request, timeout=timeout)


def _retry_after(headers: Message | None) -> float:
    """Seconds a 429 or 503 asks to wait, from Retry-After as seconds or an HTTP date, capped."""
    value = (headers.get("Retry-After", "") if headers else "").strip()
    if value.isascii() and value.isdigit():
        seconds = float(value)
    else:
        try:
            seconds = (parsedate_to_datetime(value) - datetime.now(UTC)).total_seconds()
        except (TypeError, ValueError):  # absent, malformed, or a date with no time zone
            return 0.0
    return min(max(seconds, 0.0), _MAX_RETRY_AFTER_S)


def _content_length(headers: Message) -> int | None:
    """The Content-Length header as a byte count, or None when it is absent or not a plain number."""
    value = headers.get("Content-Length", "").strip()
    return int(value) if value.isascii() and value.isdigit() else None


@dataclasses.dataclass(frozen=True, slots=True)
class _Download:
    """A body streamed to a temp file, not yet moved into place."""

    path: Path
    sha256: str
    size: int
    headers: Message
    url: str


def _download(response: _Response, source: Source, deadline: float, folder: Path) -> tuple[Path, str, int]:
    """Stream the body to a new temp file in `folder`, hashing it on the way: its path, SHA-256 and size.

    Memory stays at one chunk whatever the file's size. Refused past the source's
    `max_bytes` or past the fetch's time budget; the temp file is removed then.
    """
    encoding = response.headers.get("Content-Encoding", "identity").strip().lower()
    if encoding not in {"", "identity"}:
        msg = f"{source['id']}: served with Content-Encoding {encoding}, so its bytes are not the file itself"
        raise ValueError(msg)
    fd, temp = tempfile.mkstemp(dir=folder, prefix=f".{source['id']}.", suffix=".part")
    try:
        with os.fdopen(fd, "wb") as fh:
            digest, size = _copy(response, fh, source, deadline)
    except BaseException:
        Path(temp).unlink(missing_ok=True)
        raise
    return Path(temp), digest, size


def _copy(response: _Response, fh: IO[bytes], source: Source, deadline: float) -> tuple[str, int]:
    """Copy the body into `fh` chunk by chunk, flushed to disk: its SHA-256 and size."""
    cap = source.get("max_bytes", _MAX_BYTES)
    digest = hashlib.sha256()
    size = 0
    # read1: one socket read at most, so a server trickling bytes can't hold a read past the deadline.
    while True:
        chunk = response.read1(_CHUNK_BYTES)
        _time_left(source, deadline)
        if not chunk:
            break
        size += len(chunk)
        if size > cap:
            msg = f"{source['id']}: download exceeds {cap} bytes; raise max_bytes in sources.toml if expected"
            raise ValueError(msg)
        digest.update(chunk)
        fh.write(chunk)
    fh.flush()
    os.fsync(fh.fileno())
    return digest.hexdigest(), size


def _time_left(source: Source, deadline: float) -> float:
    """The fetch budget left for this attempt or read, or a timeout when it is spent."""
    remaining = deadline - time.monotonic()
    if remaining <= 0:
        msg = f"download still unfinished after {source.get('timeout_s', _BUDGET_S)} s"
        raise TimeoutError(msg)
    return remaining


def _get(source: Source, previous: FetchRecord | None, user_agent: str, folder: Path) -> _Download | FetchRecord:
    """GET one source's URL into a temp file in `folder`, retrying failures a publisher recovers from.

    The request is conditional on the ETag and Last-Modified of a `previous` record, which
    is returned as it is when the server answers 304 Not Modified.

    429, 500, 502, 503 and 504, network errors and a body shorter than its Content-Length
    are retried with exponential backoff, waiting at least the server's Retry-After (up to
    120 s). The whole fetch stops after `timeout_s` (default 300 s), give or take one
    socket timeout. Any other error status is raised at once.

    Raises ValueError for a body over the source's `max_bytes` (default 200 MiB), or one
    served compressed.
    """
    conditional = _conditional(previous)
    headers = {"User-Agent": user_agent, "Accept-Encoding": "identity", **_env_headers(source), **conditional}
    request = urllib.request.Request(source["url"], headers=headers)  # noqa: S310  # see _open
    deadline = time.monotonic() + source.get("timeout_s", _BUDGET_S)
    for attempt in range(1, _ATTEMPTS + 1):
        wait = _BACKOFF_S * 2 ** (attempt - 1)
        try:
            with _open(request, timeout=min(_SOCKET_TIMEOUT_S, _time_left(source, deadline))) as response:
                path, digest, size = _download(response, source, deadline, folder)
                # A connection dropped mid-body can end the read early without an error.
                expected = _content_length(response.headers)
                if expected is not None and expected != size:
                    path.unlink()
                    msg = f"body ended after {size} of the {expected} bytes its Content-Length promised"
                    raise http.client.HTTPException(msg)
                return _Download(path, digest, size, response.headers, response.url)
        except _InwardRedirectError as error:
            msg = f"{source['id']}: {error}"
            raise ValueError(msg) from error
        except urllib.error.HTTPError as error:
            # urllib raises a 304 like an error. Only a conditional request can get one.
            if error.code == HTTPStatus.NOT_MODIFIED and previous and conditional:
                return previous
            if attempt == _ATTEMPTS or error.code not in _RETRIED:
                raise
            failure: Exception = error
            wait = max(wait, _retry_after(error.headers))
        except _TRANSIENT as error:
            if attempt == _ATTEMPTS:
                raise
            failure = error
        if time.monotonic() + wait >= deadline:
            raise failure
        _logger.warning(
            "%s: attempt %d of %d failed (%s); retrying in %.0f s", source["id"], attempt, _ATTEMPTS, failure, wait
        )
        time.sleep(wait)
    raise AssertionError  # unreachable: the last attempt returns or raises


def _archive(source: Source, user_agent: str) -> str:
    """Ask the Internet Archive to snapshot the source's URL: the snapshot's URL, or "" when there is none.

    Never raises. An archive that is down, slow or answers oddly costs a warning, not the fetch.
    """
    url = source["url"]
    if urlsplit(url).scheme not in {"http", "https"}:
        _logger.warning("%s: archive needs an http or https URL, not %s; skipped", source["id"], url)
        return ""
    request = urllib.request.Request(_ARCHIVE + url, headers={"User-Agent": user_agent})  # noqa: S310  # a fixed https URL
    try:
        with _open(request, timeout=_ARCHIVE_TIMEOUT_S) as response:
            # The archive redirects to the snapshot, or names it in Content-Location.
            snapshot = urljoin(response.url, response.headers.get("Content-Location", ""))
    except (*_TRANSIENT, ValueError) as error:
        _logger.warning("%s: the Internet Archive did not snapshot %s (%s)", source["id"], url, error)
        return ""
    if not snapshot.startswith(_SNAPSHOT):
        _logger.warning("%s: the Internet Archive answered without a snapshot of %s", source["id"], url)
        return ""
    return snapshot


def _env_headers(source: Source) -> dict[str, str]:
    """The manifest's extra request headers, each value read from the environment variable it names."""
    resolved = {}
    for name, ref in source.get("headers", {}).items():
        var = ref.removeprefix("env:")
        if (value := os.environ.get(var)) is None:
            msg = f"{source['id']}: environment variable {var} for the {name} header is not set"
            raise ValueError(msg)
        resolved[name] = value
    return resolved


def _previous(source: Source, raw: Path) -> FetchRecord | None:
    """The last fetch record of `source`, if it and its raw file still match the manifest."""
    try:
        return _current(source, raw)
    except _UNUSABLE:
        return None


def _conditional(previous: FetchRecord | None) -> dict[str, str]:
    """The headers that ask the server to send the source only if it changed since `previous`."""
    if previous is None:
        return {}
    headers = {"If-None-Match": previous.etag, "If-Modified-Since": previous.last_modified}
    return {name: value for name, value in headers.items() if value}


def _is_record(fields: object) -> bool:
    """Whether `fields` holds a fetch record's id, URL, time, SHA-256 and size, each of its own type."""
    return (
        isinstance(fields, dict)
        and all(isinstance(fields.get(key), str) for key in ("source_id", "url", "retrieved_at", "sha256"))
        and _SHA256.fullmatch(fields["sha256"]) is not None
        and type(fields.get("bytes")) is int  # not isinstance: bool is an int subclass
    )


def _last_records(log: Path) -> dict[str, dict[str, Any]]:
    """The record most recently logged for each source id, from the fetch log.

    A last line cut off by an interrupted fetch is skipped with a warning. Any other line
    that is not a record is an error that names the line.
    """
    if not log.exists():
        return {}
    # Bytes that are not UTF-8 become U+FFFD, so a damaged line is reported by its number below.
    text = log.read_bytes().decode("utf-8", errors="replace")
    lines = text.splitlines()
    records: dict[str, dict[str, Any]] = {}
    numbers: dict[str, int] = {}
    for number, line in enumerate(lines, start=1):
        if not line.strip():
            continue
        try:
            record = json.loads(line)
        except json.JSONDecodeError as error:
            if number == len(lines) and not text.endswith("\n"):
                _logger.warning("%s: skipping a last line cut off by an interrupted fetch", log)
                continue
            msg = f"{log}:{number} is not a fetch record ({error})"
            raise ValueError(msg) from error
        if not (isinstance(record, dict) and isinstance(record.get("source_id"), str)):
            msg = f"{log}:{number} is not a fetch record"
            raise ValueError(msg)  # noqa: TRY004  # a damaged log: the CLI reports ValueError as one line
        records[record["source_id"]] = record
        numbers[record["source_id"]] = number
    # Only the last record per source is used, so only those get the full check: it costs
    # more than reading the line, and a log has many lines per source.
    for sid, record in records.items():
        if not _is_record(record):
            msg = f"{log}:{numbers[sid]} is not a fetch record"
            raise ValueError(msg)
    return records


def _append(log: Path, record: FetchRecord) -> None:
    """Add `record` as one line at the end of the log, in one write, flushed to disk.

    A last line cut off by an interrupted fetch is dropped first, so the new line starts
    on a line of its own. A whole record that only lacks its newline, as an editor or a
    merge can leave it, is kept. Only the end of the log is read, however long it is.

    Raises ValueError when the log is a symlink, which a cloned checkout could point at
    another file.
    """
    if log.is_symlink():
        msg = f"{log} is a symlink; the fetch log must be a file of its own"
        raise ValueError(msg)
    with log.open("a+b") as fh:
        end = fh.seek(0, os.SEEK_END)
        lead = b""
        if end and (start := _last_line_start(fh, end)) < end:
            fh.seek(start)
            try:
                json.loads(fh.read())
            except ValueError:  # not JSON, or not UTF-8
                _logger.warning("%s: dropping a last line cut off by an interrupted fetch", log)
                fh.truncate(start)
            else:
                lead = b"\n"
        fh.write(lead + json.dumps(dataclasses.asdict(record)).encode() + b"\n")
        fh.flush()
        os.fsync(fh.fileno())


def _last_line_start(fh: IO[bytes], end: int) -> int:
    """Where the line after the last newline starts: `end` when the file ends with one."""
    pos = end
    while pos > 0:
        step = min(pos, _CHUNK_BYTES)
        fh.seek(pos - step)
        if (newline := fh.read(step).rfind(b"\n")) != -1:
            return pos - step + newline + 1
        pos -= step
    return 0


def logged_versions(log: Path = FETCH_LOG) -> dict[str, str]:
    """The SHA-256 most recently logged for each source id."""
    return {sid: r["sha256"] for sid, r in _last_records(log).items()}


@contextlib.contextmanager
def _lock(raw: Path) -> Generator[None]:
    """Hold an exclusive lock on `raw/.lock`, so two fetches at once can't interleave their writes.

    Fails at once rather than waiting. POSIX only: on Windows no lock is taken.
    """
    if sys.platform == "win32":  # noqa: PLR2004
        yield
        return
    # Never truncated, and never through a symlink: a cloned checkout could plant one to another file.
    fd = os.open(raw / ".lock", os.O_WRONLY | os.O_CREAT | os.O_NOFOLLOW, 0o644)
    with os.fdopen(fd, "wb") as fh:
        try:
            fcntl.flock(fh, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError as error:
            msg = f"another fetch is already running in {raw}"
            raise OSError(msg) from error
        yield


def revisions(before: Path, *, log: Path = FETCH_LOG) -> list[str]:
    """One Markdown line per source whose bytes changed between `before`, an earlier copy of the log, and `log`.

    A source that `before` never logged is new, not revised, and is left out. Raises
    FileNotFoundError when `before` or `log` does not exist: an empty report must mean no
    revisions, not a missing copy or a run from the wrong directory.
    """
    for path in (before, log):
        if not path.exists():
            raise FileNotFoundError(errno.ENOENT, os.strerror(errno.ENOENT), str(path))
    old, new = _last_records(before), _last_records(log)
    return [
        f"- `{sid}`: {old[sid]['sha256'][:12]} -> {r['sha256'][:12]}, {r['bytes']:,} bytes, "
        f"fetched {r['retrieved_at']} from <{r['url']}>"
        for sid, r in new.items()
        if sid in old and old[sid]["sha256"] != r["sha256"]
    ]


def fetch(
    sources: list[Source],
    *,
    raw: Path = RAW,
    log: Path = FETCH_LOG,
    user_agent: str = USER_AGENT,
    only: Collection[str] | None = None,
) -> list[FetchRecord]:
    """Download each implemented source into `raw`, with a fetch record beside it.

    A source the server reports unchanged since its last intact fetch keeps that raw file
    and record, and its old record is returned. A record whose bytes differ from the last
    ones logged for its source is appended to `log`.

    Every source is tried. Raises FetchError after the loop when any failed, listing each
    with its reason; nothing is written for those. A source fails for a download over its
    `max_bytes` (default 200 MiB), bytes that miss the `sha256` the manifest pins, or an
    error that outlasts its retries. Raises OSError at once when another fetch holds the
    lock on `raw`.

    `only` restricts the fetch to those source ids, and leaves the others as they are.
    Raises ValueError for an id the manifest does not have.
    """
    sources = _select(sources, only)
    raw.mkdir(parents=True, exist_ok=True)
    log.parent.mkdir(parents=True, exist_ok=True)
    records: list[FetchRecord] = []
    failures: list[str] = []
    with _lock(raw):
        logged = _last_records(log)
        for source in implemented(sources):
            try:
                records.append(_fetch_one(source, raw, log, logged, user_agent))
            except ValueError as error:  # messages already name the source
                failures.append(str(error))
            except _TRANSIENT as error:
                failures.append(f"{source['id']}: fetching {source['url']} failed: {error}")
    if failures:
        raise FetchError(records, failures)
    return records


def _fetch_one(source: Source, raw: Path, log: Path, logged: dict[str, dict[str, Any]], user_agent: str) -> FetchRecord:
    """Fetch one source, write its raw file, log and record, and return the record."""
    response = _get(source, _previous(source, raw), user_agent, raw)
    if isinstance(response, FetchRecord):
        # A version written to disk by a fetch that died before logging it gets logged now.
        _log_new_version(log, logged, response)
        return response
    try:
        return _keep(source, response, raw, log, logged, user_agent)
    finally:
        response.path.unlink(missing_ok=True)  # gone already once moved into place


def _keep(
    source: Source, download: _Download, raw: Path, log: Path, logged: dict[str, dict[str, Any]], user_agent: str
) -> FetchRecord:
    """Check a finished download against the pin, move it into place, and log and record it.

    Bytes the log has not had for this source are first archived, when the source asks for it.
    """
    digest = download.sha256
    if source.get("sha256", digest) != digest:
        msg = f"{source['id']}: downloaded bytes hash to {digest}, not the sha256 in sources.toml"
        raise ValueError(msg)
    archived = ""
    if source.get("archive"):
        # The same bytes again, with no 304 to say so, keep the snapshot their log line names.
        last = logged.get(source["id"], {})
        archived = last.get("archive_url", "") if last.get("sha256") == digest else _archive(source, user_agent)
    record = FetchRecord(
        source_id=source["id"],
        url=source["url"],
        retrieved_at=datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ"),
        sha256=digest,
        bytes=download.size,
        vintage=_vintage(source),
        # Kept for a conditional request (If-Modified-Since, If-None-Match) on a later run.
        last_modified=download.headers.get("Last-Modified", ""),
        etag=download.headers.get("ETag", ""),
        final_url="" if download.url == source["url"] else download.url,
        env_headers=sorted(source.get("headers", {})),
        archive_url=archived,
    )
    name = raw_name(source)
    download.path.replace(raw / name)
    # Logged before the record is written: until the record matches the raw file, the next
    # run fetches in full, so a crash in between loses nothing.
    _log_new_version(log, logged, record)
    _write_atomic(raw / f"{name}.fetch.json", (json.dumps(dataclasses.asdict(record), indent=2) + "\n").encode())
    return record


def _log_new_version(log: Path, logged: dict[str, dict[str, Any]], record: FetchRecord) -> None:
    # Appended, not rewritten: the log is history. A re-download of the same bytes adds nothing.
    if logged.get(record.source_id, {}).get("sha256") != record.sha256:
        _append(log, record)
        logged[record.source_id] = dataclasses.asdict(record)


def verified(source: Source, *, raw: Path = RAW) -> tuple[FetchRecord, bytes]:
    """The fetch record and raw bytes for `source`, after checking that both are current.

    Raises ValueError when the record was made for a different request than the manifest
    now makes (another URL, vintage or `sha256` pin), or when the raw file no longer hashes
    to the record.
    """
    record = _record(source, raw)
    payload = (raw / raw_name(source)).read_bytes()
    _check_hash(source, record, hashlib.sha256(payload).hexdigest())
    return record, payload


def _current(source: Source, raw: Path) -> FetchRecord:
    """The fetch record for `source`, checked as `verified` checks it, without reading the raw file into memory."""
    record = _record(source, raw)
    _check_hash(source, record, _digest(raw / raw_name(source)))
    return record


def _digest(path: Path) -> str:
    """The file's SHA-256, read a block at a time."""
    with path.open("rb") as fh:
        return _sha256(fh)


def _sha256(fh: IO[bytes]) -> str:
    # Not hashlib.file_digest: its fixed 256 KiB buffer costs more than the hash on a small file.
    digest = hashlib.sha256()
    while chunk := fh.read(_HASH_BLOCK_BYTES):
        digest.update(chunk)
    return digest.hexdigest()


def _check_hash(source: Source, record: FetchRecord, digest: str) -> None:
    if digest != record.sha256:
        msg = f"{raw_name(source)} changed since it was fetched; re-run fetch"
        raise ValueError(msg)


def _record(source: Source, raw: Path) -> FetchRecord:
    """The fetch record for `source`, if it was made for the request the manifest now makes."""
    name = raw_name(source)
    fields = json.loads((raw / f"{name}.fetch.json").read_text(encoding="utf-8"))
    msg = f"{name}.fetch.json is not a fetch record; re-run fetch"
    if not _is_record(fields):
        raise ValueError(msg)
    try:
        record = FetchRecord(**fields)
    except TypeError as error:  # a key FetchRecord does not have, or lacks
        raise ValueError(msg) from error
    if (record.source_id, record.url) != (source["id"], source["url"]):
        msg = f"{name} was fetched from {record.url}, but sources.toml now requests {source['url']}; re-run fetch"
        raise ValueError(msg)
    if record.vintage != _vintage(source):
        msg = f"{name} was fetched as vintage {record.vintage!r}, not {_vintage(source)!r}; re-run fetch"
        raise ValueError(msg)
    if source.get("sha256", record.sha256) != record.sha256:  # load_manifest lowercases the pin
        msg = f"{name} hashes to {record.sha256}, not the sha256 sources.toml now pins; re-run fetch"
        raise ValueError(msg)
    return record


def parse_source(source: Source, payload: bytes) -> tuple[list[str], list[dict[str, str]]]:
    """One raw file as its columns and rows, read by its format: the default `Parser`.

    The format is the source's `format` key, or else the URL's extension: `.json` is JSON,
    `.xlsx` is Excel, anything else is CSV (gzipped or not). Raises ValueError for
    `format = "file"`, which has no table; `build` copies these files without a parser.
    """
    if _format(source) == FILE_FORMAT:
        msg = f"{source['id']}: format file has no table to parse"
        raise ValueError(msg)
    return _PARSERS[_format(source)](source, payload)


def _format(source: Source) -> str:
    suffix = PurePosixPath(urlsplit(source.get("url", "")).path).suffix.lower()
    return source.get("format") or {".json": "json", ".xlsx": "xlsx"}.get(suffix, "csv")


def parse_csv(source: Source, payload: bytes) -> tuple[list[str], list[dict[str, str]]]:
    """One raw CSV file as its columns and rows. A gzipped file is decompressed first.

    Raises ValueError for a repeated column name, a row with more or fewer fields than the
    header, or a gzipped file that is damaged or decompresses past the source's `max_bytes`.
    """
    payload = _maybe_gunzip(source, payload)
    # utf-8-sig drops the byte-order mark that spreadsheet exports start with. Plain UTF-8 reads the same.
    text = payload.decode("utf-8-sig")
    columns, rows = _csv_rows(source, io.StringIO(text, newline=""))
    return columns, list(rows)


def _csv_rows(source: Source, text: IO[str]) -> tuple[list[str], Iterator[dict[str, str]]]:
    """A CSV text's columns and its rows, each checked as it is read: parse_csv's checks and messages."""
    reader = csv.DictReader(text)
    columns = list(reader.fieldnames or [])
    _check_columns(source, columns)
    return columns, _checked_rows(source, reader)


def _checked_rows(source: Source, reader: csv.DictReader[str]) -> Iterator[dict[str, str]]:
    for row in reader:
        # DictReader files the surplus fields of a long row under the key None.
        if None in row:
            msg = f"{source['id']}: line {reader.line_num} has more fields than the header"
            raise ValueError(msg)
        # It fills the missing fields of a short row with None, which would publish as "".
        if None in row.values():
            msg = f"{source['id']}: line {reader.line_num} has fewer fields than the header"
            raise ValueError(msg)
        yield row


def _maybe_gunzip(source: Source, payload: bytes) -> bytes:
    """The payload decompressed when it is gzipped, else as it is.

    Refused past the source's `max_bytes`, since a small file can expand without limit,
    and when the gzip is cut off or damaged.
    """
    if payload[:2] != _GZIP_MAGIC:
        return payload
    cap = source.get("max_bytes", _MAX_BYTES)
    try:
        with gzip.GzipFile(fileobj=io.BytesIO(payload)) as fh:
            data = fh.read(cap + 1)
    except (EOFError, zlib.error, gzip.BadGzipFile) as error:
        msg = f"{source['id']}: not a readable gzip file ({error})"
        raise ValueError(msg) from error
    if len(data) > cap:
        msg = f"{source['id']}: decompresses past {cap} bytes; raise max_bytes in sources.toml if expected"
        raise ValueError(msg)
    return data


def _check_columns(source: Source, columns: list[str]) -> None:
    # A dict keeps only the last of two same-named columns, silently dropping the other.
    if dupes := sorted(c for c, n in Counter(columns).items() if n > 1):
        msg = f"{source['id']}: columns {dupes} appear more than once in the header"
        raise ValueError(msg)


class _Number(str):
    """A JSON number kept as its text, so 1.10 stays 1.10 and 1e5 stays 1e5."""

    __slots__ = ()


def _cell(value: object) -> str:
    """A JSON or spreadsheet value as output text: empty for none, `true`/`false`, JSON numbers as written."""
    match value:
        case None:
            return ""
        case str():
            return value
        case datetime() | date():
            return value.isoformat()
        case _:  # booleans as true/false, nested values as JSON text
            return _json_text(value)


def _json_text(value: object) -> str:
    """`value` as JSON text, with each `_Number` written as it was read."""
    match value:
        case _Number():
            return value
        case dict():
            return (
                "{" + ", ".join(f"{json.dumps(k, ensure_ascii=False)}: {_json_text(v)}" for k, v in value.items()) + "}"
            )
        case list():
            return "[" + ", ".join(_json_text(item) for item in value) + "]"
        case _:
            return json.dumps(value, ensure_ascii=False, default=str)


def parse_json(source: Source, payload: bytes) -> tuple[list[str], list[dict[str, str]]]:
    """A JSON array of objects as columns and rows: one column per key, in the order keys first appear.

    A key an object lacks is an empty cell. Numbers are kept as written, and nested values
    as JSON text. Raises ValueError for anything but an array of objects, or for one
    nested too deeply to read.
    """
    payload = _maybe_gunzip(source, payload)
    # From 3.14 the reader's depth limit follows the stack size, so a deep value can load and
    # then overflow while its cell is written: both steps are covered.
    try:
        data = json.loads(payload.decode("utf-8-sig"), parse_float=_Number, parse_int=_Number)
        if not (isinstance(data, list) and all(isinstance(item, dict) for item in data)):
            msg = f"{source['id']}: expected a JSON array of objects"
            raise ValueError(msg)
        columns = list(dict.fromkeys(key for item in data for key in item))
        rows = [{column: _cell(item.get(column)) for column in columns} for item in data]
    except RecursionError as error:
        msg = f"{source['id']}: JSON nested too deeply to read"
        raise ValueError(msg) from error
    return columns, rows


def parse_xlsx(source: Source, payload: bytes) -> tuple[list[str], list[dict[str, str]]]:
    """One sheet of an Excel workbook as columns and rows, the first row as the header.

    The sheet is the source's `sheet` key, or else the first. Fully empty rows are skipped,
    and so are empty header cells at the end of the row. An empty cell reads as "", like an
    empty CSV field. Needs the `xlsx` extra: `uv add 'sourceledger[xlsx]'`.

    Raises ValueError for a file that is not a workbook, one that expands past the source's
    `max_bytes`, a missing sheet, a repeated column name, or a cell right of the header.
    """
    try:
        import openpyxl  # noqa: PLC0415  # an optional extra, imported only when a source needs it
        from openpyxl.utils.exceptions import InvalidFileException  # noqa: PLC0415
    except ImportError as error:
        msg = f"{source['id']}: reading .xlsx needs the xlsx extra: uv add 'sourceledger[xlsx]'"
        raise ValueError(msg) from error
    cap = source.get("max_bytes", _MAX_BYTES)
    try:
        # The sizes the zip declares bound what it can expand to: reading past one fails its check.
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            expanded = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as error:
        msg = f"{source['id']}: not a readable Excel workbook ({error})"
        raise ValueError(msg) from error
    if expanded > cap:
        msg = f"{source['id']}: decompresses past {cap} bytes; raise max_bytes in sources.toml if expected"
        raise ValueError(msg)
    try:
        rows = _sheet_rows(source, openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True))
    except (zipfile.BadZipFile, KeyError, InvalidFileException) as error:  # a missing part reads as KeyError
        msg = f"{source['id']}: not a readable Excel workbook ({error})"
        raise ValueError(msg) from error
    if not rows:
        return [], []
    header = rows[0][1]
    width = len(header)
    while width and header[width - 1] is None:
        width -= 1
    columns = [_cell(cell) for cell in header[:width]]
    _check_columns(source, columns)
    table = []
    for number, row in rows[1:]:
        if any(cell is not None for cell in row[width:]):
            msg = f"{source['id']}: row {number} has more cells than the header"
            raise ValueError(msg)
        table.append({column: _cell(cell) for column, cell in itertools.zip_longest(columns, row[:width])})
    return columns, table


def _sheet_rows(source: Source, workbook: Any) -> list[tuple[int, tuple[object, ...]]]:  # noqa: ANN401  # openpyxl is optional
    """The sheet's rows that are not fully empty, each with its row number, and the workbook closed."""
    try:
        if (name := source.get("sheet")) is not None and name not in workbook.sheetnames:
            msg = f"{source['id']}: no sheet {name!r}; the workbook has {workbook.sheetnames}"
            raise ValueError(msg)
        sheet = workbook[name] if name is not None else workbook.worksheets[0]
        return [
            (number, row)
            for number, row in enumerate(sheet.iter_rows(values_only=True), start=1)
            if any(cell is not None for cell in row)
        ]
    finally:
        workbook.close()


_PARSERS: dict[str, Parser] = {"csv": parse_csv, "json": parse_json, "xlsx": parse_xlsx}


def _stamp(record: FetchRecord) -> dict[str, str]:
    """The provenance columns `build` adds to every output row."""
    return {
        "source_id": record.source_id,
        "source_url": record.url,
        "retrieved_at": record.retrieved_at,
        "vintage": record.vintage,
        "raw_sha256": record.sha256,
    }


def build(
    sources: list[Source],
    *,
    raw: Path = RAW,
    processed: Path = PROCESSED,
    parse: Parser = parse_source,
    only: Collection[str] | None = None,
) -> list[Path]:
    """Write implemented sources to `processed`: unchanged files, or tables stamped with their fetch records.

    `format = "file"` copies the original bytes to `processed/<raw_name>`, without parsing.
    For tables, `parse` reads each raw file; the default reads CSV, JSON and Excel by format.
    Then writes `processed/datapackage.json`, a Frictionless Data Package (v2) describing every
    output: its columns, size and SHA-256, its source URL and licence, and its fetch record.
    The same records give the same file, byte for byte.

    `only` restricts the build to those source ids. The package keeps each other source's
    entry as the last build wrote it. Raises ValueError for an id the manifest does not have,
    or a file output named `datapackage.json`, which is reserved for the descriptor.
    """
    selected = {s["id"] for s in _select(sources, only)}
    earlier = {} if only is None else _listed(processed / DATAPACKAGE)
    processed.mkdir(parents=True, exist_ok=True)
    written = []
    resources = []
    for source in implemented(sources):
        if source["id"] not in selected:
            if source["id"] in earlier:
                resources.append(earlier[source["id"]])
            continue
        out = processed / _output_name(source)
        record, fields = _build_one(source, raw, out, parse)
        written.append(out)
        resources.append(_resource(source, record, fields, out))
    package = {"$schema": _DATAPACKAGE_PROFILE, "name": _package_name(), "resources": resources}
    _write_atomic(processed / DATAPACKAGE, (json.dumps(package, indent=2, ensure_ascii=False) + "\n").encode())
    return written


def _build_one(source: Source, raw: Path, out: Path, parse: Parser) -> tuple[FetchRecord, list[str]]:
    """Write one source to `out`: the fetch record, and the columns written (none for a file).

    File sources are copied unchanged. Tables get their provenance columns first.
    With the default parser a plain CSV file goes through a row at a time, so memory stays
    near one row whatever the file's size.
    """
    record = _record(source, raw)
    stamp = _stamp(record)
    with (raw / raw_name(source)).open("rb") as fh:
        _check_hash(source, record, _sha256(fh))
        fh.seek(0)
        if source.get("format") == FILE_FORMAT:
            with _replacing(out) as target:
                shutil.copyfileobj(fh, target)
            return record, []
        if parse is parse_source and _PARSERS[_format(source)] is parse_csv and fh.read(2) != _GZIP_MAGIC:
            fh.seek(0)
            text = io.TextIOWrapper(fh, encoding="utf-8-sig", newline="")
            try:
                columns, rows = _csv_rows(source, text)
                _write_rows(source, out, stamp, columns, rows)
            except (ValueError, csv.Error):
                # parse_csv decodes the whole file before it reads a row, so it can fail on another
                # line, or name another byte offset. Reading the file whole gives exactly its error.
                fh.seek(0)
                parse(source, fh.read())
                raise
            finally:
                text.detach()  # closing the wrapper would close fh
        else:
            # NOTE: other formats, gzipped CSV and custom parsers hold the whole file and its rows
            # in memory, as a Parser returns lists; stream them too once such sources outgrow memory.
            fh.seek(0)
            columns, rows = parse(source, fh.read())
            _write_rows(source, out, stamp, columns, rows)
    return record, [*stamp, *columns]


def _write_rows(
    source: Source, out: Path, stamp: dict[str, str], columns: list[str], rows: Iterable[dict[str, str]]
) -> None:
    """Write `rows` to `out` as CSV, each with the provenance columns first, as it gets them."""
    if clash := stamp.keys() & set(columns):
        msg = f"{source['id']}: source columns {sorted(clash)} collide with the provenance columns"
        raise ValueError(msg)
    with _replacing(out) as fh:
        text = io.TextIOWrapper(fh, encoding="utf-8", newline="")
        writer = csv.DictWriter(text, fieldnames=[*stamp, *columns])
        writer.writeheader()
        # The stamp last: a parser's row cannot overwrite a provenance column.
        writer.writerows({**row, **stamp} for row in rows)
        text.detach()  # flushes; _replacing closes fh


def _listed(path: Path) -> dict[str, Any]:
    """The resources a data package lists, by name: none when it is missing or unreadable."""
    try:
        package = json.loads(path.read_text(encoding="utf-8"))
    except _UNUSABLE:
        return {}
    return _by_name(package)


def _by_name(package: object) -> dict[str, Any]:
    """A parsed data package's resources, by name: none when it is not an object with a list of them."""
    resources = package.get("resources") if isinstance(package, dict) else None
    return {r.get("name"): r for r in resources if isinstance(r, dict)} if isinstance(resources, list) else {}


def _package_name() -> str:
    """The project directory's name, lower-cased, as Data Package names should be."""
    return re.sub(r"[^a-z0-9._-]+", "-", Path.cwd().name.lower()).strip("-") or "dataset"


def _resource(source: Source, record: FetchRecord, fields: list[str], out: Path) -> dict[str, Any]:
    """One output as a Data Package resource: a file, or a table whose columns hold text."""
    resource: dict[str, Any] = {
        "name": source["id"],
        "type": "table",
        "path": out.name,
        "format": "csv",
        "mediatype": "text/csv",
        "encoding": "utf-8",
        "bytes": out.stat().st_size,
        "hash": f"sha256:{_digest(out)}",
        "schema": {"fields": [{"name": field, "type": "string"} for field in fields]},
        "sources": [{"title": str(source.get("name", source["id"])), "path": record.url}],
    }
    if source.get("format") == FILE_FORMAT:
        for key in ("type", "format", "mediatype", "encoding", "schema"):
            del resource[key]
    if source.get("license"):
        licence = {"name": source["license"], "path": source.get("terms_url"), "title": source.get("terms")}
        resource["licenses"] = [{k: v for k, v in licence.items() if v}]
    # Custom properties are allowed in Data Package v2; these name the bytes the rows came from.
    resource["sourceledger"] = {
        "retrieved_at": record.retrieved_at,
        "vintage": record.vintage,
        "raw_sha256": record.sha256,
        "final_url": record.final_url,
        "redistribute": source.get("redistribute"),
    }
    return resource


def read_output(path: Path) -> list[dict[str, str]]:
    """A published output file, as written."""
    with path.open(newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def _terms_problems(source: Source) -> list[str]:
    """What an implemented source is missing of its reuse terms: in the publisher's words, and machine-readable."""
    sid, problems = source["id"], []
    if not (source.get("terms") and source.get("terms_url")):
        problems.append(f"{sid}: no terms or terms_url; record the publisher's reuse terms")
    license_ = source.get("license")
    if not (isinstance(license_, str) and license_.strip()):
        problems.append(f"{sid}: no license; give the SPDX id of the terms, or LicenseRef-<name> for custom ones")
    if not isinstance(source.get("redistribute"), bool):
        problems.append(f"{sid}: redistribute must be true or false: may the raw file be republished?")
    return problems


def validate(
    sources: list[Source], *, raw: Path = RAW, processed: Path = PROCESSED, only: Collection[str] | None = None
) -> list[str]:
    """Check the published output against the manifest and the fetch records, and return one string per problem.

    `only` restricts the checks to those source ids, and skips the two that concern the whole
    output: files no source owns, and the package's list of sources. Raises ValueError for an
    id the manifest does not have.
    """
    selected = _select(sources, only)
    problems = []
    outputs = {_output_name(source) for source in sources}
    # Output of a source since removed or renamed would otherwise pass as published data.
    if only is None:
        problems += [
            f"{out}: no source {out.stem!r} in sources.toml; delete it"
            for out in sorted(processed.glob("*"))
            if out.is_file() and out.name != DATAPACKAGE and out.name not in outputs
        ]
    for source in selected:
        sid, status = source["id"], source.get("status")
        out = processed / _output_name(source)
        if status == UNIMPLEMENTED:
            if out.exists():
                problems.append(f"{sid}: status is unimplemented, but {out} exists")
            continue
        if status != IMPLEMENTED:
            problems.append(f"{sid}: status must be {IMPLEMENTED!r} or {UNIMPLEMENTED!r}, not {status!r}")
            continue
        problems += _terms_problems(source)
        problems += _output_problems(source, raw, out)
    return problems + _package_problems(sources, processed, only)


def _output_problems(source: Source, raw: Path, out: Path) -> list[str]:
    """Whether one built file still matches the current fetch."""
    try:
        record = _current(source, raw)
        if source.get("format") == FILE_FORMAT:
            return [f"{out}: differs from the fetched file; re-run build"] if _digest(out) != record.sha256 else []
        rows, stale = _scan(out, _stamp(record))
    except (OSError, ValueError, csv.Error) as error:
        return [one_line(error)]
    problems = []
    if not rows:
        problems.append(f"{out}: no rows")
    if stale:
        problems.append(f"{out}: rows come from an earlier fetch than {raw_name(source)}; re-run build")
    return problems


def _scan(out: Path, stamp: dict[str, str]) -> tuple[bool, bool]:
    """Whether an output has rows, and whether any of them has another stamp, read one row at a time."""
    rows = stale = False
    with out.open(newline="", encoding="utf-8") as fh:
        # Read to the end even once stale, so a damaged file fails as csv.Error wherever it is damaged.
        for row in csv.DictReader(fh):
            rows = True
            stale = stale or {column: row.get(column) for column in stamp} != stamp
    return rows, stale


def _package_problems(sources: list[Source], processed: Path, only: Collection[str] | None) -> list[str]:
    """Whether datapackage.json lists exactly the implemented sources, each with its output's current hash.

    With `only`, whether it lists those of them, each with its output's current hash.
    """
    path = processed / DATAPACKAGE
    expected = {source["id"] for source in implemented(sources)}
    checked = expected if only is None else expected & set(only)
    try:
        package = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        return [f"{path} is missing; re-run build"] if checked else []
    except ValueError as error:
        return [f"{path} is not JSON ({error}); re-run build"]
    listed = _by_name(package)
    if only is None and set(listed) != expected:
        return [
            f"{path} lists {sorted(map(str, listed))}, not the implemented sources {sorted(expected)}; re-run build"
        ]
    problems = []
    by_id = {source["id"]: source for source in sources}
    for sid in sorted(checked):
        source = by_id[sid]
        out = processed / _output_name(source)
        if sid not in listed:
            problems.append(f"{path} does not list {sid}; re-run build")
        elif out.exists() and listed[sid].get("hash") != f"sha256:{_digest(out)}":
            problems.append(f"{path} describes an earlier build of {out}; re-run build")
    return problems


def _yes_no(value: object) -> str:
    """`redistribute` as yes or no, or `?` when it is neither true nor false."""
    if not isinstance(value, bool):
        return "?"
    return "yes" if value else "no"


def status(sources: list[Source], *, log: Path = FETCH_LOG) -> list[str]:
    """One line per source, from the manifest and the fetch log only: no network, no raw files.

    Each line gives the id, the status, the vintage, when the version last logged was fetched
    (or `never`), the start of its SHA-256, the licence, and whether the raw file may be
    republished.
    """
    last = _last_records(log)
    width = max((len(s["id"]) for s in sources), default=0)
    lines = []
    for source in sources:
        record = last.get(source["id"], {})
        fields = [
            f"{source['id']:<{width}}",
            f"{source.get('status')!s:<13}",
            _vintage(source),
            record.get("retrieved_at", "never"),
            record.get("sha256", "-")[:12],
            str(source.get("license") or "-"),
            f"redistribute={_yes_no(source.get('redistribute'))}",
        ]
        lines.append("  ".join(fields))
    return lines


def _md_cell(value: object) -> str:
    """A value as one Markdown table cell: on one line, with its pipes escaped."""
    return " ".join(str(value).split()).replace("|", "\\|")


def attribution(sources: list[Source]) -> list[str]:
    """The implemented sources as the lines of a Markdown table, in manifest order.

    One row per source: its name (or id), the terms in the publisher's words, the licence,
    where the terms are stated, and whether the raw file may be republished. For a
    dataset's README or release notes.
    """
    lines = ["| Source | Terms | Licence | Terms URL | Redistribute |", "| --- | --- | --- | --- | --- |"]
    for source in implemented(sources):
        values = [
            source.get("name", source["id"]),
            source.get("terms", ""),
            source.get("license", ""),
            source.get("terms_url", ""),
            _yes_no(source.get("redistribute")),
        ]
        lines.append("| " + " | ".join(map(_md_cell, values)) + " |")
    return lines


def verify(folder: Path) -> list[str]:
    """Check a published dataset against its datapackage.json, offline, and return one string per problem.

    Each resource's file must be inside `folder` and match the `bytes` (when given) and the
    `hash` (`sha256:<hex>`) the package lists. Raises OSError when the package is missing,
    and ValueError when it is not JSON.
    """
    path = folder / DATAPACKAGE
    package = json.loads(path.read_text(encoding="utf-8"))
    resources = package.get("resources") if isinstance(package, dict) else None
    if not isinstance(resources, list):
        return [f"{path} has no list of resources"]
    return [problem for r in resources for problem in _file_problems(folder, r if isinstance(r, dict) else {})]


def _file_problems(folder: Path, resource: dict[str, Any]) -> list[str]:
    """Whether one resource's file is in `folder`, with the size and SHA-256 the package lists."""
    name = resource.get("path")
    # The package comes from someone else, so a path may not point outside the folder.
    if not isinstance(name, str) or urlsplit(name).scheme or Path(name).is_absolute() or os.pardir in Path(name).parts:
        return [f"{folder / DATAPACKAGE}: resource path {name!r} is not a file in {folder}"]
    file = folder / name
    if not file.is_file():
        return [f"{file} is missing"]
    # A symlink inside the folder can still lead out of it.
    if not file.resolve().is_relative_to(folder.resolve()):
        return [f"{folder / DATAPACKAGE}: resource path {name!r} is not a file in {folder}"]
    problems = []
    size = file.stat().st_size
    # Not isinstance: bool is an int subclass, so `"bytes": true` would pass for a 1-byte file.
    if (listed := resource.get("bytes")) is not None and (type(listed) is not int or listed != size):
        shown = f"{listed:,}" if type(listed) is int else repr(listed)  # from someone else's file
        problems.append(f"{file} is {size:,} bytes, not the {shown} the package lists")
    expected = resource.get("hash")
    if not (isinstance(expected, str) and expected.lower().startswith("sha256:")):
        problems.append(f"{file}: the package lists no sha256 hash to check it against")
        return problems
    digest = _digest(file)
    if f"sha256:{digest}" != expected.lower():
        problems.append(f"{file} hashes to sha256:{digest[:12]}, not the hash the package lists")
    return problems


def one_line(error: Exception) -> str:
    """One line for an expected failure: a missing file by name, anything else as raised."""
    if isinstance(error, FileNotFoundError) and error.filename:
        return f"{error.filename} is missing"
    return str(error)
